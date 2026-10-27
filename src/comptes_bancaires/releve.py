#!/usr/bin/env python3
# -*- coding: utf-8 -*-
r"""
Relevé de comptes bancaires (exports CSV BoursoBank) -> objets Python -> Excel.

Utilisation rapide (ligne de commande, après installation) :
    comptes donnees\exports\*.csv                     # ajoute au classeur donnees\Comptes.xlsx
    comptes export_marion.csv --compte Marion
    comptes donnees\exports\*.csv --nouveau           # repart d'un classeur vierge
    comptes export.csv --analyse Nourriture            # coût par contexte
    comptes --importer-config donnees\Comptes.xlsx     # onglets Excel -> JSON
    comptes-graphiques                                 # courbes interactives (HTML)

Par défaut, la configuration, le classeur et les graphiques sont dans le dossier
« donnees » du répertoire courant (voir chemin_config_par_defaut).

Utilisation en Python :
    from comptes_bancaires import Config, lire_csv, calculer_totaux, ecrire_excel, analyser_contextes
    cfg = Config.charger("config_comptes.json")
    ops = lire_csv("export-operations.csv", cfg)
    calculer_totaux(ops, cfg)
    ecrire_excel(ops, "Comptes.xlsx", cfg)
    print(analyser_contextes(ops, cfg, "Nourriture"))

Toutes les listes (types, moyens de paiement, correspondances de libellés,
numéros de compte, contextes) sont dans config_comptes.json et peuvent être
enrichies à la main ou avec Config.ajouter_type / ajouter_moyen_paiement /
ajouter_correspondance / associer_compte / ajouter_contexte.

CONTEXTES
---------
Un contexte (vacances, mariage, déménagement...) a une date de début et de fin.
Chaque opération reçoit le nom du contexte où elle tombe, sinon « Quotidien ».
Le type ne change pas : on peut donc croiser Type × Contexte (ex. Nourriture en
vacances vs au quotidien, ramené à un coût par jour). Un contexte peut en plus :
  - avoir ses propres règles regex, prioritaires sur les règles globales ;
  - remplacer des types (ex. {"*": "Mariage"} : tout ce qui est payé pendant
    le contexte passe en type Mariage) ;
  - ne concerner que certains comptes.
Pour les paiements par carte, c'est la date d'ACHAT (« CARTE 29/09/26 » dans le
libellé) qui sert à trouver le contexte, pas la date de débit : un repas du
dimanche débité le mardi reste bien dans les vacances.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Optional

import glob
import os
import shutil
import sys

DOSSIER_DONNEES = "donnees"
NOM_CONFIG = "config_comptes.json"


def chemin_config_par_defaut() -> Path:
    """
    Ordre de recherche : variable d'environnement COMPTES_CONFIG, puis
    <racine>/donnees/config_comptes.json, puis <racine>/config_comptes.json, où <racine> est
    COMPTES_RACINE (fixée par les lanceurs py comptes.py...) ou le dossier courant. Si rien n'existe,
    copie la configuration modèle fournie avec le paquet dans ./donnees/.
    """
    if os.environ.get("COMPTES_CONFIG"):
        return Path(os.environ["COMPTES_CONFIG"])
    base = Path(os.environ.get("COMPTES_RACINE", "."))     # fixé par les lanceurs py comptes.py...
    for c in (base / DOSSIER_DONNEES / NOM_CONFIG, base / NOM_CONFIG):
        if c.exists():
            return c
    cible = base / DOSSIER_DONNEES / NOM_CONFIG
    cible.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(Path(__file__).with_name("config_defaut.json"), cible)
    print(f"Configuration modèle créée : {cible}")
    return cible


def charger_config(chemin: str | Path | None = None) -> "Config":
    """Charge la configuration (défaut : chemin_config_par_defaut) en la migrant si elle est ancienne."""
    from .migration import migrer_config
    chemin = Path(chemin) if chemin else chemin_config_par_defaut()
    journal = migrer_config(chemin)
    if journal:
        print(f"Configuration {chemin} mise à jour :")
        for ligne in journal:
            print(f"  - {ligne}")
    return Config.charger(chemin)


def etendre_motifs(motifs: list[str]) -> list[str]:
    """Développe *.csv (l'invite de commandes Windows ne le fait pas elle-même)."""
    res = []
    for m in motifs:
        trouves = sorted(glob.glob(m)) if any(ch in m for ch in "*?[") else [m]
        if not trouves:
            raise SystemExit(f"Aucun fichier ne correspond à : {m}")
        res += trouves
    return res


def console_utf8() -> None:
    """Évite les erreurs d'accents / € dans une console Windows redirigée."""
    for flux in (sys.stdout, sys.stderr):
        try:
            flux.reconfigure(errors="replace")
        except (AttributeError, ValueError):
            pass


def sauver_classeur(wb, chemin: Path) -> None:
    try:
        wb.save(chemin)
    except PermissionError:
        raise SystemExit(f"Impossible d'écrire {chemin} : le fichier est sans doute ouvert dans "
                         f"Excel. Fermez-le puis relancez la commande.")


MOIS = ["Janvier", "Février", "Mars", "Avril", "Mai", "Juin", "Juillet",
        "Août", "Septembre", "Octobre", "Novembre", "Décembre"]
QUOTIDIEN = "Quotidien"


def _date(v) -> date:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(str(v).strip(), fmt).date()
        except ValueError:
            pass
    raise ValueError(f"Date non reconnue : {v!r} (attendu AAAA-MM-JJ ou JJ/MM/AAAA)")


def texte_date(d: date) -> str:
    return f"{d.day} {MOIS[d.month - 1]} {d.year}"


# --------------------------------------------------------------------------- #
#  Contexte
# --------------------------------------------------------------------------- #
@dataclass
class Contexte:
    nom: str
    debut: date
    fin: date
    categorie: str = ""                      # ex. "Vacances", "Événement" (pour regrouper)
    comptes: Optional[list[str]] = None      # None = tous les comptes
    regles: list[dict] = field(default_factory=list)          # comme "libelles", prioritaires
    remplacer_types: dict[str, str] = field(default_factory=dict)  # {"Loisirs": "Voyage"} ; "*" = tous
    actif: bool = True

    def __post_init__(self):
        self.debut, self.fin = _date(self.debut), _date(self.fin)
        if self.fin < self.debut:
            raise ValueError(f"Contexte {self.nom!r} : la fin est avant le début")
        self._re = [(re.compile(r["regex"], re.I), r["libelle"], r.get("type")) for r in self.regles]

    @classmethod
    def depuis_dict(cls, d: dict) -> "Contexte":
        return cls(nom=d["nom"], debut=d["debut"], fin=d["fin"], categorie=d.get("categorie", ""),
                   comptes=d.get("comptes"), regles=d.get("regles", []),
                   remplacer_types=d.get("remplacer_types", {}), actif=d.get("actif", True))

    def vers_dict(self) -> dict:
        return {"nom": self.nom, "debut": self.debut.isoformat(), "fin": self.fin.isoformat(),
                "categorie": self.categorie, "comptes": self.comptes, "regles": self.regles,
                "remplacer_types": self.remplacer_types, "actif": self.actif}

    @property
    def nb_jours(self) -> int:
        return (self.fin - self.debut).days + 1

    def contient(self, d: date, compte: Optional[str] = None) -> bool:
        if not self.actif or not (self.debut <= d <= self.fin):
            return False
        return self.comptes is None or compte is None or compte in self.comptes

    def appliquer_type(self, type_: str) -> str:
        return self.remplacer_types.get(type_, self.remplacer_types.get("*", type_))

    def __str__(self) -> str:
        return f"{self.nom} ({texte_date(self.debut)} → {texte_date(self.fin)}, {self.nb_jours} j)"


# --------------------------------------------------------------------------- #
#  Configuration (bases de données éditables)
# --------------------------------------------------------------------------- #
@dataclass
class Config:
    chemin: Path
    data: dict

    @classmethod
    def charger(cls, chemin: str | Path = "config_comptes.json") -> "Config":
        chemin = Path(chemin)
        with open(chemin, encoding="utf-8") as f:
            cfg = cls(chemin, json.load(f))
        cfg.data.setdefault("contextes", [])
        cfg._compiler()
        return cfg

    def _compiler(self) -> None:
        self._re_moyens = [(re.compile(r["regex"], re.I), r["moyen"])
                           for r in self.data.get("moyens_paiement", [])]
        self._re_libelles = [(re.compile(r["regex"], re.I), r["libelle"], r.get("type"))
                             for r in self.data.get("libelles", [])]
        self.contextes = [Contexte.depuis_dict(c) for c in self.data.get("contextes", [])]

    def sauvegarder(self) -> None:
        anciens = {d.get("nom"): d for d in self.data.get("contextes", [])}
        self.data["contextes"] = [{**anciens.get(c.nom, {}), **c.vers_dict()} for c in self.contextes]
        with open(self.chemin, "w", encoding="utf-8") as f:
            json.dump(self.data, f, ensure_ascii=False, indent=2)
        self._compiler()

    # ---- accès -------------------------------------------------------------
    @property
    def types(self) -> list[str]:
        return self.data["types"]

    @property
    def noms_comptes(self) -> list[str]:
        return self.data["noms_comptes"]

    @property
    def liste_moyens(self) -> list[str]:
        vus = []
        for r in self.data.get("moyens_paiement", []):
            if r["moyen"] not in vus:
                vus.append(r["moyen"])
        defaut = self.data.get("moyen_par_defaut", "Autre")
        if defaut not in vus:
            vus.append(defaut)
        return vus

    @property
    def noms_contextes(self) -> list[str]:
        return [QUOTIDIEN] + [c.nom for c in self.contextes if c.actif]

    def contexte(self, nom: str) -> Optional[Contexte]:
        return next((c for c in self.contextes if c.nom == nom), None)

    def contexte_pour(self, d: date, compte: Optional[str] = None) -> Optional[Contexte]:
        """Premier contexte (dans l'ordre de la config) qui contient la date."""
        return next((c for c in self.contextes if c.contient(d, compte)), None)

    # ---- enrichissement ----------------------------------------------------
    def ajouter_type(self, nom: str, sauver: bool = True) -> None:
        if nom not in self.types:
            self.types.append(nom)
            if sauver:
                self.sauvegarder()

    def ajouter_moyen_paiement(self, regex: str, moyen: str, sauver: bool = True) -> None:
        self.data["moyens_paiement"].append({"regex": regex, "moyen": moyen})
        if sauver:
            self.sauvegarder()

    def ajouter_correspondance(self, regex: str, libelle: str,
                               type_: Optional[str] = None, en_tete: bool = True,
                               sauver: bool = True) -> None:
        """Ajoute une règle regex -> libellé (et type). en_tete=True : prioritaire."""
        re.compile(regex)  # lève une erreur tout de suite si la regex est invalide
        if type_ and type_ not in self.types:
            self.ajouter_type(type_, sauver=False)
        regle = {"regex": regex, "libelle": libelle, "type": type_}
        if en_tete:
            self.data["libelles"].insert(0, regle)
        else:
            self.data["libelles"].append(regle)
        if sauver:
            self.sauvegarder()

    def associer_compte(self, numero: str, nom: str, sauver: bool = True) -> None:
        self.data["comptes"][numero] = nom
        if nom not in self.noms_comptes:
            self.noms_comptes.append(nom)
        if sauver:
            self.sauvegarder()

    def ajouter_contexte(self, nom: str, debut, fin, categorie: str = "",
                         comptes: Optional[list[str]] = None,
                         regles: Optional[list[dict]] = None,
                         remplacer_types: Optional[dict[str, str]] = None,
                         sauver: bool = True) -> Contexte:
        """
        Ex. : cfg.ajouter_contexte("Vacances Corse", "2026-08-01", "2026-08-15", "Vacances")
              cfg.ajouter_contexte("Mariage Julie", "2026-09-12", "2026-09-13", "Événement",
                                   remplacer_types={"*": "Mariage"})
        """
        if self.contexte(nom):
            raise ValueError(f"Le contexte {nom!r} existe déjà")
        c = Contexte(nom, debut, fin, categorie, comptes, regles or [], remplacer_types or {})
        for t in c.remplacer_types.values():
            self.ajouter_type(t, sauver=False)
        self.contextes.append(c)
        if sauver:
            self.sauvegarder()
        return c

    # ---- règles ------------------------------------------------------------
    def moyen_paiement(self, libelle_brut: str) -> str:
        for rx, moyen in self._re_moyens:
            if rx.search(libelle_brut):
                return moyen
        return self.data.get("moyen_par_defaut", "Autre")

    def libelle_et_type(self, libelle_brut: str, libelle_suggere: str,
                        categorie: str, categorie_parente: str,
                        ctx: Optional[Contexte] = None) -> tuple[str, str]:
        libelle, type_ = None, None
        regles = (ctx._re if ctx else []) + self._re_libelles   # règles du contexte d'abord
        for rx, lib, typ in regles:
            if rx.search(libelle_brut):
                libelle, type_ = lib, typ
                break
        if libelle is None:
            libelle = libelle_suggere.strip() or nettoyer_libelle(libelle_brut)
        if not type_:
            cats = self.data.get("categories_banque", {})
            type_ = cats.get(categorie) or cats.get(categorie_parente) \
                or self.data.get("type_par_defaut", "À classer")
        if ctx:
            type_ = ctx.appliquer_type(type_)
        return libelle, type_


def nettoyer_libelle(brut: str) -> str:
    """Retire 'CARTE 29/09/26', 'CB*7901', 'PRLV SEPA'... pour un libellé lisible."""
    s = re.sub(r"^(CARTE|RETRAIT DAB)\s+\d{2}/\d{2}/\d{2}\s+", "", brut, flags=re.I)
    s = re.sub(r"\s+CB\*\d+$", "", s, flags=re.I)
    s = re.sub(r"^(PRLV SEPA|VIR SEPA|VIR INST|VIR)\s+", "", s, flags=re.I)
    return s.strip().title()


def date_achat(libelle_brut: str, date_operation: date) -> date:
    """Date réelle d'un paiement carte / retrait (dans le libellé), sinon date d'opération."""
    m = re.match(r"^(?:CARTE|RETRAIT DAB)\s+(\d{2})/(\d{2})/(\d{2})\b", libelle_brut, re.I)
    if m:
        try:
            return date(2000 + int(m.group(3)), int(m.group(2)), int(m.group(1)))
        except ValueError:
            pass
    return date_operation


# --------------------------------------------------------------------------- #
#  Objet Opération
# --------------------------------------------------------------------------- #
@dataclass
class Operation:
    date: date
    compte: str                 # Cyrille / Marion / Joint
    type: str                   # Nourriture, Santé, ...
    libelle: str                # libellé choisi (via regex)
    debit: float = 0.0          # > 0 si dépense
    credit: float = 0.0         # > 0 si revenu
    moyen_paiement: str = ""
    pointage: bool = False
    contexte: str = QUOTIDIEN   # nom du contexte (vacances...) ou "Quotidien"
    commentaire: str = ""
    libelle_banque: str = ""    # libellé brut du CSV (pour contrôle / doublons)
    date_achat: Optional[date] = None
    solde_banque: Optional[float] = None
    totaux: dict = field(default_factory=dict)   # {"Cyrille": x, "Marion": y, "Joint": z, "Total": t}

    @property
    def date_texte(self) -> str:
        """Ex. : '30 Septembre 2026'."""
        return texte_date(self.date)

    @property
    def montant(self) -> float:
        return round(self.credit - self.debit, 2)

    @property
    def total_cyrille(self): return self.totaux.get("Cyrille")
    @property
    def total_marion(self): return self.totaux.get("Marion")
    @property
    def total_joint(self): return self.totaux.get("Joint")
    @property
    def total_general(self): return self.totaux.get("Total")

    def cle(self) -> tuple:
        """Clé utilisée pour éviter les doublons lors d'un ajout au fichier Excel."""
        return (self.date, self.compte, self.libelle_banque.strip().upper(),
                round(self.debit, 2), round(self.credit, 2))

    def __str__(self) -> str:
        sens = f"-{self.debit:.2f} €" if self.debit else f"+{self.credit:.2f} €"
        return (f"{self.date_texte:<20} {self.compte:<8} {self.type:<14} {self.contexte:<16} "
                f"{self.libelle:<28} {sens:>12}  {self.moyen_paiement}")


# --------------------------------------------------------------------------- #
#  Lecture du CSV
# --------------------------------------------------------------------------- #
def _nombre(txt: str) -> float:
    """'-1 800,00' / '413.98' / '2 736,36' -> float."""
    s = (txt or "").strip().replace(" ", "").replace(" ", "").replace(" ", "")
    if not s:
        return 0.0
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    return float(s)


def _ouvrir(chemin: Path) -> list[list[str]]:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            with open(chemin, encoding=enc, newline="") as f:
                return list(csv.reader(f, delimiter=";"))
        except UnicodeDecodeError:
            continue
    raise ValueError(f"Encodage non reconnu : {chemin}")


def lire_csv(chemin: str | Path, cfg: Config, compte: Optional[str] = None) -> list[Operation]:
    """
    Lit un export CSV BoursoBank et renvoie la liste des opérations,
    triées de la plus ancienne à la plus récente.
    `compte` force le nom du compte (sinon déduit du 'Numéro Compte' via la config).
    """
    lignes = [l for l in _ouvrir(Path(chemin)) if any(c.strip() for c in l)]
    entete = [h.strip() for h in lignes[0]]

    def idx(nom: str, dernier: bool = False, requis: bool = True) -> Optional[int]:
        positions = [i for i, h in enumerate(entete) if h == nom]
        if not positions:
            if requis:
                raise KeyError(f"Colonne '{nom}' absente du CSV ({entete})")
            return None
        return positions[-1] if dernier else positions[0]

    # Le CSV Bourso contient deux colonnes 'Solde' : la 1re = montant, la 2e = solde du compte
    i_date, i_lib, i_sugg = idx("Date Opération"), idx("Libellé"), idx("Libellé Suggéré")
    i_cat, i_catp = idx("Catégorie"), idx("Catégorie Parente")
    i_mont, i_solde = idx("Solde"), idx("Solde", dernier=True)
    i_num, i_point = idx("Numéro Compte"), idx("Pointage")
    i_com = idx("Commentaire", requis=False)

    ops: list[Operation] = []
    for l in lignes[1:]:
        l = l + [""] * (len(entete) - len(l))
        brut = l[i_lib].strip()
        montant = _nombre(l[i_mont])
        nom_compte = compte or cfg.data["comptes"].get(l[i_num].strip())
        if not nom_compte:
            raise ValueError(
                f"Numéro de compte {l[i_num]!r} inconnu : ajoutez-le dans 'comptes' "
                f"de la config, ou utilisez compte='Marion' / --compte Marion.")
        d_op = _date(l[i_date])
        d_achat = date_achat(brut, d_op)
        ctx = cfg.contexte_pour(d_achat, nom_compte)
        libelle, type_ = cfg.libelle_et_type(brut, l[i_sugg], l[i_cat].strip(),
                                             l[i_catp].strip(), ctx)
        ops.append(Operation(
            date=d_op,
            compte=nom_compte,
            type=type_,
            libelle=libelle,
            debit=round(-montant, 2) if montant < 0 else 0.0,
            credit=round(montant, 2) if montant > 0 else 0.0,
            moyen_paiement=cfg.moyen_paiement(brut),
            pointage=l[i_point].strip().lower() in ("oui", "o", "vrai", "true", "1", "x"),
            contexte=ctx.nom if ctx else QUOTIDIEN,
            commentaire=l[i_com].strip() if i_com is not None else "",
            libelle_banque=brut,
            date_achat=d_achat,
            solde_banque=_nombre(l[i_solde]) if l[i_solde].strip() else None,
        ))
    # Le CSV est du plus récent au plus ancien : on inverse (tri stable par date ensuite)
    ops.reverse()
    ops.sort(key=lambda o: o.date)
    return ops


def appliquer_contextes(ops: list[Operation], cfg: Config) -> None:
    """
    (Ré)attribue le contexte de chaque opération d'après les dates de la config.
    Une opération placée à la main dans un contexte hors de ses dates est conservée.
    """
    for o in ops:
        ctx = cfg.contexte_pour(o.date_achat or o.date, o.compte)
        if ctx:
            o.contexte = ctx.nom
        elif o.contexte not in (cfg.noms_contextes + [""]) or not o.contexte:
            o.contexte = QUOTIDIEN


def soldes_initiaux_depuis_csv(ops: list[Operation]) -> dict[str, float]:
    """
    Solde de départ de chaque compte = solde affiché par la banque à la date la plus
    récente - somme de toutes les opérations du fichier.
    """
    res = {}
    for nom in {o.compte for o in ops}:
        ops_c = [o for o in ops if o.compte == nom]
        date_max = max(o.date for o in ops_c)
        soldes = [o.solde_banque for o in ops_c if o.date == date_max and o.solde_banque is not None]
        if soldes:
            res[nom] = round(soldes[-1] - sum(o.montant for o in ops_c), 2)
    return res


def calculer_totaux(ops: list[Operation], cfg: Config,
                    soldes_initiaux: Optional[dict[str, float]] = None) -> dict[str, float]:
    """Renseigne op.totaux (solde de chaque compte + total général après l'opération)."""
    soldes = {nom: 0.0 for nom in cfg.noms_comptes}
    soldes.update(soldes_initiaux if soldes_initiaux is not None else soldes_initiaux_depuis_csv(ops))
    depart = dict(soldes)
    for o in sorted(ops, key=lambda o: o.date):
        soldes[o.compte] = round(soldes.get(o.compte, 0.0) + o.montant, 2)
        o.totaux = dict(soldes)
        o.totaux["Total"] = round(sum(soldes.values()), 2)
    return depart


# --------------------------------------------------------------------------- #
#  Analyse par contexte
# --------------------------------------------------------------------------- #
def analyser_contextes(ops: list[Operation], cfg: Config, type_: Optional[str] = None,
                       comptes: Optional[list[str]] = None) -> list[dict]:
    """
    Dépenses (débits - remboursements) par contexte, totales et ramenées par jour.
    type_ : ex. "Nourriture" (None = toutes les dépenses, hors "types_hors_analyse"
            de la config : virements entre vous, charges fixes...).
    La ligne « Quotidien » compte les jours de la période qui ne sont dans aucun contexte.
    """
    exclus = set(cfg.data.get("types_hors_analyse", []))
    sel = [o for o in ops if (o.type == type_ if type_ else o.type not in exclus)
           and (comptes is None or o.compte in comptes)]
    if not ops:
        return []
    d0 = min(o.date for o in ops)          # même période que dans Excel (dates d'opération)
    d1 = max(o.date for o in ops)
    jours_ctx = set()
    lignes = []
    for c in cfg.contextes:
        if not c.actif:
            continue
        deb, fin = max(c.debut, d0), min(c.fin, d1)
        nb = max(0, (fin - deb).days + 1)
        jours_ctx |= {deb + timedelta(i) for i in range(nb)}
        dep = sum(o.debit - o.credit for o in sel if o.contexte == c.nom)
        lignes.append({"contexte": c.nom, "categorie": c.categorie, "debut": c.debut, "fin": c.fin,
                       "jours": nb, "depense": round(dep, 2),
                       "par_jour": round(dep / nb, 2) if nb else None})
    nb_q = (d1 - d0).days + 1 - len(jours_ctx)
    dep_q = sum(o.debit - o.credit for o in sel if o.contexte == QUOTIDIEN)
    lignes.insert(0, {"contexte": QUOTIDIEN, "categorie": "", "debut": d0, "fin": d1,
                      "jours": nb_q, "depense": round(dep_q, 2),
                      "par_jour": round(dep_q / nb_q, 2) if nb_q else None})
    ref = lignes[0]["par_jour"]
    for l in lignes:
        l["ecart_vs_quotidien"] = (round(l["par_jour"] - ref, 2)
                                   if l["par_jour"] is not None and ref is not None else None)
    return lignes


def afficher_analyse(lignes: list[dict], titre: str) -> None:
    print(f"\n=== {titre} ===")
    print(f"{'Contexte':<26}{'Jours':>6}{'Dépense':>12}{'€/jour':>10}{'Écart/jour':>12}")
    for l in lignes:
        pj = f"{l['par_jour']:.2f}" if l["par_jour"] is not None else "-"
        ec = f"{l['ecart_vs_quotidien']:+.2f}" if l["ecart_vs_quotidien"] is not None else "-"
        print(f"{l['contexte']:<26}{l['jours']:>6}{l['depense']:>12.2f}{pj:>10}{ec:>12}")


# --------------------------------------------------------------------------- #
#  Écriture Excel
# --------------------------------------------------------------------------- #
FMT_EUR = '#,##0.00 €;-#,##0.00 €;""'
FMT_EUR0 = '#,##0.00 €;-#,##0.00 €;0.00 €'
FMT_DATE = '[$-fr-FR]d mmmm yyyy'


def _lire_existant(chemin: Path, cfg: Config):
    """Relit un classeur déjà produit : opérations (avec vos corrections) + soldes initiaux."""
    from openpyxl import load_workbook
    from openpyxl.utils.datetime import from_excel
    wb = load_workbook(chemin)
    ops, soldes = [], {}
    if "Paramètres" in wb.sheetnames:
        p = wb["Paramètres"]
        for r in range(2, p.max_row + 1):
            nom, val = p.cell(r, 1).value, p.cell(r, 2).value
            if nom and isinstance(val, (int, float)):
                soldes[str(nom)] = float(val)
        for r in range(2, p.max_row + 1):        # types ajoutés directement dans Excel
            v = p.cell(r, 4).value
            if v and str(v) not in cfg.types:
                cfg.types.append(str(v))
    ws = wb["Opérations"]
    entete = [c.value for c in ws[1]]
    col = {nom: i for i, nom in enumerate(entete) if nom}

    def val(row, nom, defaut=""):
        return row[col[nom]] if nom in col and row[col[nom]] is not None else defaut

    for row in ws.iter_rows(min_row=2, values_only=True):
        d = row[col["Date"]]
        if isinstance(d, (int, float)) and 30000 < d < 80000:   # date stockée en numéro de série
            d = from_excel(d)
        if not isinstance(d, (datetime, date)):
            continue
        d = _date(d)
        brut = str(val(row, "Libellé banque"))
        ops.append(Operation(
            date=d, compte=val(row, "Compte"), type=val(row, "Type"),
            libelle=val(row, "Libellé"),
            debit=float(val(row, "Débit", 0) or 0), credit=float(val(row, "Crédit", 0) or 0),
            moyen_paiement=val(row, "Moyen de paiement"),
            pointage=str(val(row, "Pointage")).lower() == "oui",
            contexte=val(row, "Contexte", QUOTIDIEN), commentaire=str(val(row, "Commentaire")),
            libelle_banque=brut, date_achat=date_achat(brut, d),
        ))
    return ops, soldes


def ecrire_excel(ops: list[Operation], chemin: str | Path, cfg: Config,
                 soldes_initiaux: Optional[dict[str, float]] = None,
                 ajouter: bool = True) -> Path:
    """
    Écrit les opérations dans un classeur Excel mis en forme.
    ajouter=True : si le fichier existe, les opérations déjà présentes (et vos
    modifications manuelles de Type / Libellé / Pointage / Commentaire) sont conservées,
    seules les nouvelles sont ajoutées (doublons ignorés). Les contextes sont
    réappliqués d'après la config, pour tenir compte d'un contexte ajouté après coup.
    """
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as L
    from openpyxl.worksheet.datavalidation import DataValidation

    chemin = Path(chemin)
    soldes: dict[str, float] = {}
    existantes: list[Operation] = []
    if ajouter and chemin.exists():
        existantes, soldes = _lire_existant(chemin, cfg)

    cles = {o.cle() for o in existantes}
    nouvelles = [o for o in ops if o.cle() not in cles]
    toutes = sorted(existantes + nouvelles, key=lambda o: o.date)
    appliquer_contextes(toutes, cfg)

    calc = soldes_initiaux if soldes_initiaux is not None else soldes_initiaux_depuis_csv(nouvelles)
    for nom, v in calc.items():
        soldes.setdefault(nom, v)
    for nom in cfg.noms_comptes:
        soldes.setdefault(nom, 0.0)
    noms = list(dict.fromkeys(cfg.noms_comptes + list(soldes)))

    # ---- colonnes de la feuille Opérations ----------------------------------
    entetes = (["Date", "Compte", "Type", "Contexte", "Libellé", "Débit", "Crédit",
                "Moyen de paiement", "Pointage"]
               + [f"Total {n}" for n in noms] + ["Total général", "Commentaire", "Libellé banque"])
    C = {nom: i + 1 for i, nom in enumerate(entetes)}          # nom de colonne -> numéro
    CL = {nom: L(i) for nom, i in C.items()}                   # nom de colonne -> lettre
    c_tot0 = C[f"Total {noms[0]}"]

    # ---- styles -------------------------------------------------------------
    police = "Arial"
    f_norm, f_bold = Font(name=police, size=10), Font(name=police, size=10, bold=True)
    f_head = Font(name=police, size=10, bold=True, color="FFFFFF")
    f_input = Font(name=police, size=10, color="0000FF")
    f_titre = Font(name=police, size=12, bold=True)
    f_note = Font(name=police, size=9, italic=True, color="595959")
    fill_head = PatternFill("solid", start_color="1F3864")
    fill_jaune = PatternFill("solid", start_color="FFF2CC")
    fill_vert = PatternFill("solid", start_color="E2EFDA")
    fill_gris = PatternFill("solid", start_color="F2F2F2")
    fine = Side(style="thin", color="BFBFBF")
    bord = Border(bottom=fine)
    centre = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def entete_ligne(ws, ligne, titres, col0=1):
        for c, t in enumerate(titres, col0):
            cell = ws.cell(ligne, c, t)
            cell.font, cell.fill, cell.alignment = f_head, fill_head, centre

    wb = Workbook()
    ws = wb.active
    ws.title = "Opérations"
    wp = wb.create_sheet("Paramètres")
    wsy = wb.create_sheet("Synthèse")

    # ---- Paramètres ----------------------------------------------------------
    entete_ligne(wp, 1, ["Compte", "Solde initial"])
    entete_ligne(wp, 1, ["Types", "Moyens de paiement"], 4)
    entete_ligne(wp, 1, ["Aide"], 7)
    lignes_solde = {}
    for r, nom in enumerate(noms, 2):
        wp.cell(r, 1, nom).font = f_bold
        cell = wp.cell(r, 2, soldes.get(nom, 0.0))
        cell.number_format, cell.font, cell.fill = FMT_EUR0, f_input, fill_jaune
        lignes_solde[nom] = r
    r_tot = 2 + len(noms)
    wp.cell(r_tot, 1, "Total").font = f_bold
    cell = wp.cell(r_tot, 2, f"=SUM(B2:B{r_tot - 1})")
    cell.number_format, cell.font = FMT_EUR0, f_bold
    for r, t in enumerate(cfg.types, 2):
        wp.cell(r, 4, t).font = f_norm
    for r, m in enumerate(cfg.liste_moyens, 2):
        wp.cell(r, 5, m).font = f_norm
    aide = [
        "Cellules jaunes / texte bleu = à saisir.",
        "Solde initial : solde du compte AVANT la 1re opération",
        "(calculé automatiquement depuis le CSV : solde banque - opérations).",
        "Ajoutez des Types / Moyens en bas des colonnes D et E :",
        "ils apparaissent dans les listes déroulantes de 'Opérations'.",
        "Les onglets gris (Règles, Contextes, Types...) reprennent config_comptes.json :",
        "modifiez-les puis lancez  --importer-config Comptes.xlsx  pour mettre le JSON à jour.",
    ]
    for r, t in enumerate(aide, 2):
        wp.cell(r, 7, t).font = f_note
    for col, w in zip("ABCDEFG", (12, 15, 3, 18, 20, 3, 80)):
        wp.column_dimensions[col].width = w
    len([c for c in cfg.contextes if c.actif])

    # ---- Opérations ---------------------------------------------------------
    entete_ligne(ws, 1, entetes)
    ws.row_dimensions[1].height = 28
    for r, o in enumerate(toutes, 2):
        valeurs = {"Date": o.date, "Compte": o.compte, "Type": o.type, "Contexte": o.contexte,
                   "Libellé": o.libelle, "Débit": o.debit or None, "Crédit": o.credit or None,
                   "Moyen de paiement": o.moyen_paiement,
                   "Pointage": "Oui" if o.pointage else "Non",
                   "Commentaire": o.commentaire or None, "Libellé banque": o.libelle_banque}
        for nom, v in valeurs.items():
            cell = ws.cell(r, C[nom], v)
            cell.font, cell.border = f_norm, bord
        ws.cell(r, C["Date"]).number_format = FMT_DATE
        for nom, couleur in (("Débit", "C00000"), ("Crédit", "007A33")):
            ws.cell(r, C[nom]).number_format = FMT_EUR
            ws.cell(r, C[nom]).font = Font(name=police, size=10, color=couleur)
        ws.cell(r, C["Pointage"]).alignment = Alignment(horizontal="center")
        ws.cell(r, C["Libellé banque"]).font = Font(name=police, size=9, color="808080")
        ws.cell(r, C["Commentaire"]).font = Font(name=police, size=10, italic=True)
        for nom in noms:
            col = CL[f"Total {nom}"]
            prec = f"Paramètres!$B${lignes_solde[nom]}" if r == 2 else f"{col}{r - 1}"
            cell = ws.cell(r, C[f"Total {nom}"],
                           f'={prec}+IF(${CL["Compte"]}{r}="{nom}",'
                           f'N(${CL["Crédit"]}{r})-N(${CL["Débit"]}{r}),0)')
            cell.number_format, cell.font, cell.border = FMT_EUR0, f_norm, bord
        cell = ws.cell(r, C["Total général"],
                       f"=SUM({L(c_tot0)}{r}:{CL[f'Total {noms[-1]}']}{r})")
        cell.number_format, cell.font, cell.border = FMT_EUR0, f_bold, bord
        cell.fill = fill_gris

    der = max(len(toutes) + 1, 2)
    largeurs = {"Date": 18, "Compte": 10, "Type": 14, "Contexte": 20, "Libellé": 28,
                "Débit": 12, "Crédit": 12, "Moyen de paiement": 16, "Pointage": 9,
                "Total général": 14, "Commentaire": 32, "Libellé banque": 42}
    for nom in entetes:
        ws.column_dimensions[CL[nom]].width = largeurs.get(nom, 13)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{CL['Libellé banque']}{der}"

    fin = der + 500                       # lignes en plus pour vos saisies manuelles
    listes = [
        ("Compte", f"=Paramètres!$A$2:$A${1 + len(noms)}"),
        ("Type", "=Paramètres!$D$2:$D$200"),
        ("Moyen de paiement", "=Paramètres!$E$2:$E$100"),
        ("Pointage", '"Oui,Non"'),
    ]   # la liste "Contexte" est ajoutée plus bas (elle pointe vers le tableau de la Synthèse)
    ws.conditional_formatting.add(f"A2:{CL['Libellé banque']}{fin}",
                                  FormulaRule(formula=[f'${CL["Pointage"]}2="Oui"'], fill=fill_vert))

    # ---- Synthèse -----------------------------------------------------------
    plage = lambda nom: f"Opérations!${CL[nom]}$2:${CL[nom]}${fin}"

    # 1) Soldes
    wsy.cell(1, 1, "Soldes actuels").font = f_titre
    entete_ligne(wsy, 2, noms + ["Total général"], 2)
    for k, nom in enumerate(noms):
        cell = wsy.cell(3, 2 + k, f'=Paramètres!$B${lignes_solde[nom]}'
                                  f'+SUMIFS({plage("Crédit")},{plage("Compte")},"{nom}")'
                                  f'-SUMIFS({plage("Débit")},{plage("Compte")},"{nom}")')
        cell.number_format, cell.font = FMT_EUR0, f_bold
    cell = wsy.cell(3, 2 + len(noms), f"=SUM(B3:{L(1 + len(noms))}3)")
    cell.number_format, cell.font, cell.fill = FMT_EUR0, f_bold, fill_gris
    wsy.cell(3, 1, "Solde").font = f_bold

    # 2) Par type
    wsy.cell(5, 1, "Dépenses et revenus par type").font = f_titre
    tete = ["Type"] + [f"Débit {n}" for n in noms] + ["Total débit", "Total crédit"]
    entete_ligne(wsy, 6, tete)
    r0 = 7
    for i, t in enumerate(cfg.types):
        r = r0 + i
        wsy.cell(r, 1, t).font = f_norm
        for k, nom in enumerate(noms):
            cell = wsy.cell(r, 2 + k, f'=SUMIFS({plage("Débit")},{plage("Type")},$A{r},'
                                      f'{plage("Compte")},"{nom}")')
            cell.number_format, cell.font = FMT_EUR, f_norm
        cd = 2 + len(noms)
        cell = wsy.cell(r, cd, f"=SUM(B{r}:{L(cd - 1)}{r})")
        cell.number_format, cell.font = FMT_EUR, f_bold
        cell = wsy.cell(r, cd + 1, f'=SUMIFS({plage("Crédit")},{plage("Type")},$A{r})')
        cell.number_format, cell.font = FMT_EUR, f_bold
    rt = r0 + len(cfg.types)
    wsy.cell(rt, 1, "Total").font = f_bold
    for c in range(2, len(tete) + 1):
        cell = wsy.cell(rt, c, f"=SUM({L(c)}{r0}:{L(c)}{rt - 1})")
        cell.number_format, cell.font = FMT_EUR0, f_bold
        cell.border = Border(top=Side(style="thin"))

    # 3) Analyse par contexte (type au choix)
    ra = rt + 3
    wsy.cell(ra, 1, "Coût par contexte").font = f_titre
    wsy.cell(ra + 1, 1, "Type analysé :").font = f_bold
    cell = wsy.cell(ra + 1, 2, "Nourriture" if "Nourriture" in cfg.types else cfg.types[0])
    cell.font, cell.fill = f_input, fill_jaune
    wsy.cell(ra + 1, 3, "← choisissez un type, ou « Tous » (hors virements et charges fixes)").font = f_note
    choix = '"' + ",".join(["Tous"] + [t.replace(",", " ") for t in cfg.types]) + '"'
    dv_type = DataValidation(type="list", formula1=choix if len(choix) < 250
                             else "=Paramètres!$D$2:$D$200", showErrorMessage=False)
    dv_type.add(wsy.cell(ra + 1, 2))
    wsy.add_data_validation(dv_type)
    type_ref = f"$B${ra + 1}"

    wsy.cell(ra + 2, 1, "Période des données :").font = f_bold
    p_deb, p_fin = f"$B${ra + 2}", f"$C${ra + 2}"
    for c, f in ((2, f"=MIN({plage('Date')})"), (3, f"=MAX({plage('Date')})")):
        cell = wsy.cell(ra + 2, c, f)
        cell.number_format, cell.font = FMT_DATE, f_norm

    rh = ra + 4
    entete_ligne(wsy, rh, ["Contexte", "Début", "Fin", "Jours", "Dépense nette",
                           "€ / jour", "Écart / jour vs Quotidien", "Nb opérations"])
    wsy.row_dimensions[rh].height = 30
    critere_type = f'IF({type_ref}="Tous","*",{type_ref})'
    exclus = [t for t in cfg.data.get("types_hors_analyse", []) if t in cfg.types]

    def ligne_ctx(r, nom, deb, fin_, jours):
        wsy.cell(r, 1, nom).font = f_bold if nom == QUOTIDIEN else f_norm
        for c, v in ((2, deb), (3, fin_)):
            cell = wsy.cell(r, c, v)
            cell.number_format, cell.font = FMT_DATE, f_norm
        wsy.cell(r, 4, jours).font = f_norm
        net = lambda crit: (f'SUMIFS({plage("Débit")},{plage("Contexte")},$A{r},{plage("Type")},{crit})'
                            f'-SUMIFS({plage("Crédit")},{plage("Contexte")},$A{r},{plage("Type")},{crit})')
        dep = net(critere_type)
        if exclus:   # « Tous » = tout sauf les types hors analyse
            dep += "-IF(" + type_ref + '="Tous",' + "+".join(f'({net(chr(34) + t + chr(34))})' for t in exclus) + ",0)"
        nb = (f'COUNTIFS({plage("Contexte")},$A{r},{plage("Type")},{critere_type})'
              + ("-IF(" + type_ref + '="Tous",' + "+".join(
                  f'COUNTIFS({plage("Contexte")},$A{r},{plage("Type")},"{t}")' for t in exclus) + ",0)"
                 if exclus else ""))
        cell = wsy.cell(r, 5, f"={dep}")
        cell.number_format, cell.font = FMT_EUR0, f_norm
        cell = wsy.cell(r, 6, f'=IF(D{r}>0,E{r}/D{r},"")')
        cell.number_format, cell.font, cell.fill = FMT_EUR0, f_bold, fill_gris
        cell = wsy.cell(r, 7, f'=IF(AND(D{r}>0,$D${rh + 1}>0),F{r}-$F${rh + 1},"")')
        cell.number_format, cell.font = '+#,##0.00 €;-#,##0.00 €;0.00 €', f_norm
        cell = wsy.cell(r, 8, f"={nb}")
        cell.font = f_norm
        for c in range(1, 9):
            wsy.cell(r, c).border = bord

    actifs = [c for c in cfg.contextes if c.actif]
    r_q = rh + 1
    r_c0, r_c1 = r_q + 1, r_q + len(actifs)
    # Jours d'un contexte, limités à la période des données
    jours_ctx = lambda r: f"=MAX(0,MIN(C{r},{p_fin})-MAX(B{r},{p_deb})+1)"
    somme_ctx = f"SUM(D{r_c0}:D{r_c1})" if actifs else "0"
    ligne_ctx(r_q, QUOTIDIEN, f"={p_deb}", f"={p_fin}", f"={p_fin}-{p_deb}+1-{somme_ctx}")
    for i, c in enumerate(actifs):
        ligne_ctx(r_c0 + i, c.nom, c.debut, c.fin, jours_ctx(r_c0 + i))
    r_note = r_q + len(actifs) + 2
    for k, t in enumerate([
        "Dépense nette = débits - crédits du type choisi (un remboursement réduit la dépense).",
        "« Tous » exclut les types listés dans \"types_hors_analyse\" de la config : "
        + ", ".join(cfg.data.get("types_hors_analyse", [])) + ".",
        "Quotidien = jours de la période qui ne sont dans aucun contexte (les contextes ne doivent pas se chevaucher).",
        "Le contexte d'un paiement carte est déterminé par sa date d'achat, pas sa date de débit.",
    ]):
        wsy.cell(r_note + k, 1, t).font = f_note

    # Liste déroulante "Contexte" de la feuille Opérations = colonne A de ce tableau
    dv = DataValidation(type="list", formula1=f"=Synthèse!$A${r_q}:$A${max(r_c1, r_q)}",
                        allow_blank=True)
    dv.add(f"{CL['Contexte']}2:{CL['Contexte']}{fin}")
    ws.add_data_validation(dv)
    for nom, source in listes:
        dv = DataValidation(type="list", formula1=source, allow_blank=True)
        dv.add(f"{CL[nom]}2:{CL[nom]}{fin}")
        ws.add_data_validation(dv)

    wsy.column_dimensions["A"].width = 26
    for c in range(2, 10):
        wsy.column_dimensions[L(c)].width = 16

    # ---- Évolution mensuelle + onglets de configuration ----------------------
    _feuille_evolution(wb, cfg, toutes, CL, fin)
    _feuilles_config(wb, cfg, toutes)

    sauver_classeur(wb, chemin)
    print(f"{len(nouvelles)} opération(s) ajoutée(s), {len(toutes)} au total -> {chemin}")
    return chemin


# --------------------------------------------------------------------------- #
#  Onglets supplémentaires : évolution mensuelle et configuration
# --------------------------------------------------------------------------- #
_POLICE = "Arial"


def _styles():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    return {
        "norm": Font(name=_POLICE, size=10), "bold": Font(name=_POLICE, size=10, bold=True),
        "head": Font(name=_POLICE, size=10, bold=True, color="FFFFFF"),
        "input": Font(name=_POLICE, size=10, color="0000FF"),
        "stat": Font(name=_POLICE, size=10, color="595959"),
        "titre": Font(name=_POLICE, size=12, bold=True),
        "note": Font(name=_POLICE, size=9, italic=True, color="595959"),
        "fill_head": PatternFill("solid", start_color="1F3864"),
        "fill_cfg": PatternFill("solid", start_color="44546A"),
        "fill_gris": PatternFill("solid", start_color="F2F2F2"),
        "bord": Border(bottom=Side(style="thin", color="BFBFBF")),
        "centre": Alignment(horizontal="center", vertical="center", wrap_text=True),
    }


def _tableau(ws, entetes, lignes, largeurs, editables=(), formats=None, note=None, cfg_sheet=True):
    """Écrit un tableau simple : en-tête, lignes, largeurs. Colonnes 'editables' en bleu."""
    st = _styles()
    from openpyxl.utils import get_column_letter as L
    for c, t in enumerate(entetes, 1):
        cell = ws.cell(1, c, t)
        cell.font, cell.alignment = st["head"], st["centre"]
        cell.fill = st["fill_cfg"] if cfg_sheet else st["fill_head"]
    for r, ligne in enumerate(lignes, 2):
        for c, v in enumerate(ligne, 1):
            cell = ws.cell(r, c, v)
            cell.border = st["bord"]
            cell.font = st["input"] if entetes[c - 1] in editables else st["stat"]
            if formats and entetes[c - 1] in formats:
                cell.number_format = formats[entetes[c - 1]]
    for c, w in enumerate(largeurs, 1):
        ws.column_dimensions[L(c)].width = w
    ws.freeze_panes = "A2"
    if lignes:
        ws.auto_filter.ref = f"A1:{L(len(entetes))}{len(lignes) + 1}"
    if note:
        ws.cell(len(lignes) + 3, 1, note).font = st["note"]
    if cfg_sheet:
        ws.sheet_properties.tabColor = "8497B0"


def _liste_types(ws, colonne: str, nb_lignes: int) -> None:
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type="list", formula1="=Paramètres!$D$2:$D$200", allow_blank=True)
    dv.add(f"{colonne}2:{colonne}{nb_lignes + 200}")
    ws.add_data_validation(dv)


def _feuille_evolution(wb, cfg: Config, toutes: list[Operation], CL: dict, fin: int) -> None:
    """Tableau mensuel Type × Mois (formules) + graphique de la catégorie choisie."""
    from openpyxl.chart import LineChart, Reference
    from openpyxl.utils import get_column_letter as L
    from openpyxl.worksheet.datavalidation import DataValidation
    st = _styles()
    ws = wb.create_sheet("Évolution")
    if not toutes:
        return
    d0, d1 = min(o.date for o in toutes), max(o.date for o in toutes)
    mois, m = [], date(d0.year, d0.month, 1)
    while m <= d1:
        mois.append(m)
        m = date(m.year + (m.month == 12), m.month % 12 + 1, 1)
    exclus = set(cfg.data.get("types_hors_analyse", []))
    types = [t for t in cfg.types]
    plage = lambda nom: f"Opérations!${CL[nom]}$2:${CL[nom]}${fin}"

    ws.cell(1, 1, "Dépense nette par mois et par type (débits - crédits)").font = st["titre"]
    ws.cell(2, 1, "Catégorie du graphique :").font = st["bold"]
    sel = ws.cell(2, 3, "Nourriture" if "Nourriture" in types else types[0])
    sel.font = st["input"]
    from openpyxl.styles import PatternFill
    sel.fill = PatternFill("solid", start_color="FFF2CC")
    dv = DataValidation(type="list", formula1="=Paramètres!$D$2:$D$200")
    dv.add(sel)
    ws.add_data_validation(dv)

    h = 4                                   # ligne d'en-tête du tableau
    entetes = ["Mois"] + types + ["Total (hors virements / charges fixes)", "Catégorie choisie"]
    for c, t in enumerate(entetes, 1):
        cell = ws.cell(h, c, t)
        cell.font, cell.fill, cell.alignment = st["head"], st["fill_head"], st["centre"]
    ws.row_dimensions[h].height = 45
    c_tot, c_sel = len(types) + 2, len(types) + 3
    for i, mo in enumerate(mois):
        r = h + 1 + i
        cell = ws.cell(r, 1, mo)
        cell.number_format, cell.font = '[$-fr-FR]mmmm yyyy', st["bold"]
        crit_date = (f'{plage("Date")},">="&$A{r},{plage("Date")},"<"&EDATE($A{r},1)')
        for k, t in enumerate(types):
            ref = f"{L(2 + k)}${h}"
            f = (f'=SUMIFS({plage("Débit")},{plage("Type")},{ref},{crit_date})'
                 f'-SUMIFS({plage("Crédit")},{plage("Type")},{ref},{crit_date})')
            cell = ws.cell(r, 2 + k, f)
            cell.number_format, cell.font, cell.border = FMT_EUR, st["norm"], st["bord"]
        a_sommer = [f"{L(2 + k)}{r}" for k, t in enumerate(types) if t not in exclus]
        cell = ws.cell(r, c_tot, "=" + "+".join(a_sommer))
        cell.number_format, cell.font, cell.fill = FMT_EUR0, st["bold"], st["fill_gris"]
        cell = ws.cell(r, c_sel, f"=INDEX($B{r}:${L(len(types) + 1)}{r},"
                                 f"MATCH($C$2,$B${h}:${L(len(types) + 1)}${h},0))")
        cell.number_format, cell.font, cell.fill = FMT_EUR0, st["bold"], st["fill_gris"]
    r_der = h + len(mois)
    r_moy = r_der + 1
    ws.cell(r_moy, 1, "Moyenne / mois").font = st["bold"]
    for c in range(2, c_sel + 1):
        cell = ws.cell(r_moy, c, f"=AVERAGE({L(c)}{h + 1}:{L(c)}{r_der})")
        cell.number_format, cell.font = FMT_EUR0, st["bold"]
    ws.cell(r_moy + 2, 1, "Les mois incomplets (début / fin d'export) tirent la moyenne vers le bas. "
                          "Les crédits (remboursements, salaire) sont déduits : une valeur négative "
                          "= rentrée d'argent.").font = st["note"]
    ws.column_dimensions["A"].width = 16
    for c in range(2, c_sel + 1):
        ws.column_dimensions[L(c)].width = 13
    ws.freeze_panes = ws.cell(h + 1, 2)

    def graphique(col, titre, ancre):
        ch = LineChart()
        ch.title, ch.height, ch.width = titre, 8, 18
        ch.y_axis.title, ch.y_axis.numFmt = "€", '#,##0 €'
        ch.x_axis.number_format = 'mmm yy'
        ch.legend = None
        ch.add_data(Reference(ws, min_col=col, min_row=h, max_row=r_der), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=h + 1, max_row=r_der))
        s = ch.series[0]
        s.graphicalProperties.line.solidFill, s.graphicalProperties.line.width = "2A78D6", 22000
        s.marker.symbol, s.marker.size = "circle", 7
        s.marker.graphicalProperties.solidFill = "2A78D6"
        ch.x_axis.delete = ch.y_axis.delete = False
        ws.add_chart(ch, ancre)

    graphique(c_sel, "Catégorie choisie (case C2)", f"A{r_moy + 4}")
    graphique(c_tot, "Total des dépenses courantes", f"H{r_moy + 4}")


def _feuilles_config(wb, cfg: Config, toutes: list[Operation]) -> None:
    """Un onglet par rubrique de config_comptes.json (+ statistiques d'utilisation)."""
    _styles()
    exclus = set(cfg.data.get("types_hors_analyse", []))

    # Règles de libellés, avec le nombre d'opérations que chaque règle capte réellement
    regles = cfg.data.get("libelles", [])
    rx = [re.compile(r["regex"], re.I) for r in regles]
    stats = [[0, 0.0, ""] for _ in regles]
    sans_regle: dict[str, list] = {}
    for o in toutes:
        k = next((i for i, x in enumerate(rx) if x.search(o.libelle_banque)), None)
        if k is None:
            cle = nettoyer_libelle(o.libelle_banque) or o.libelle
            g = sans_regle.setdefault(cle, [0, 0.0, o.libelle_banque, o.type])
            g[0] += 1
            g[1] += o.debit - o.credit
        else:
            stats[k][0] += 1
            stats[k][1] += o.debit - o.credit
            stats[k][2] = stats[k][2] or o.libelle_banque
    ws = wb.create_sheet("Règles libellés")
    lignes = [[i + 1, r["regex"], r["libelle"], r.get("type") or "", n, round(m, 2), ex]
              for i, (r, (n, m, ex)) in enumerate(zip(regles, stats))]
    _tableau(ws, ["Ordre", "Regex", "Libellé", "Type", "Nb opérations", "Dépense nette",
                  "Exemple (libellé banque)"], lignes, [7, 34, 28, 15, 13, 14, 44],
             editables=("Ordre", "Regex", "Libellé", "Type"), formats={"Dépense nette": FMT_EUR0},
             note="Bleu = modifiable. Ordre = priorité (la 1re règle qui correspond gagne). "
                  "Type vide = la catégorie de la banque décide. Nb opérations = calculé.")
    _liste_types(ws, "D", len(lignes))

    # Opérations sans règle : aide à créer de nouvelles règles
    ws = wb.create_sheet("Sans règle")
    def suggestion(cle: str) -> str:
        mots = re.sub(r"^(SAS|SASU|SARL|SA|EURL)\s+|\s+\d+$", "", cle.upper()).split()
        return re.escape(" ".join(mots[:2])) if mots else ""
    lignes = sorted(([cle, n, round(m, 2), t, ex, suggestion(cle)]
                     for cle, (n, m, ex, t) in sans_regle.items()), key=lambda l: -l[1])
    _tableau(ws, ["Libellé nettoyé", "Nb opérations", "Dépense nette", "Type attribué",
                  "Exemple (libellé banque)", "Regex suggérée"], lignes,
             [30, 13, 14, 15, 44, 22], formats={"Dépense nette": FMT_EUR0},
             note="Opérations qu'aucune règle ne capte (le type vient alors de la catégorie banque). "
                  "Copiez une ligne dans 'Règles libellés' pour la fixer.")

    # Contextes et leurs règles propres
    ws = wb.create_sheet("Contextes")
    lignes = []
    for r, c in enumerate(cfg.contextes, 2):
        lignes.append([c.nom, c.categorie, c.debut, c.fin, f"=D{r}-C{r}+1",
                       ", ".join(c.comptes) if c.comptes else "Tous",
                       "; ".join(f"{k} → {v}" for k, v in c.remplacer_types.items()),
                       "Oui" if c.actif else "Non",
                       sum(1 for o in toutes if o.contexte == c.nom),
                       round(sum(o.debit - o.credit for o in toutes
                                 if o.contexte == c.nom and o.type not in exclus), 2)])
    _tableau(ws, ["Nom", "Catégorie", "Début", "Fin", "Jours", "Comptes", "Types remplacés",
                  "Actif", "Nb opérations", "Dépenses courantes"], lignes,
             [26, 13, 18, 18, 7, 16, 26, 7, 13, 16],
             editables=("Nom", "Catégorie", "Début", "Fin", "Comptes", "Types remplacés", "Actif"),
             formats={"Début": FMT_DATE, "Fin": FMT_DATE, "Dépenses courantes": FMT_EUR0},
             note="Comptes : « Tous » ou « Cyrille, Joint ». Types remplacés : « Loisirs → Voyage; "
                  "* → Mariage » (* = tous). Règles propres : onglet 'Règles contextes'.")

    ws = wb.create_sheet("Règles contextes")
    lignes = [[c.nom, r["regex"], r["libelle"], r.get("type") or ""]
              for c in cfg.contextes for r in c.regles]
    _tableau(ws, ["Contexte", "Regex", "Libellé", "Type"], lignes, [26, 34, 28, 15],
             editables=("Contexte", "Regex", "Libellé", "Type"),
             note="Règles appliquées seulement pendant le contexte, avant les règles globales.")
    _liste_types(ws, "D", len(lignes))

    # Types
    ws = wb.create_sheet("Types")
    lignes = []
    for t in cfg.types:
        sel = [o for o in toutes if o.type == t]
        lignes.append([t, "Oui" if t in exclus else "Non", len(sel),
                       round(sum(o.debit for o in sel), 2), round(sum(o.credit for o in sel), 2)])
    _tableau(ws, ["Type", "Hors analyse", "Nb opérations", "Débits", "Crédits"], lignes,
             [18, 13, 13, 14, 14], editables=("Type", "Hors analyse"),
             formats={"Débits": FMT_EUR0, "Crédits": FMT_EUR0},
             note="Hors analyse = Oui : exclu des totaux « Tous » (virements entre vous, charges fixes...).")

    # Moyens de paiement
    ws = wb.create_sheet("Moyens paiement")
    lignes = [[i + 1, r["regex"], r["moyen"],
               sum(1 for o in toutes if o.moyen_paiement == r["moyen"])]
              for i, r in enumerate(cfg.data.get("moyens_paiement", []))]
    lignes.append(["défaut", "", cfg.data.get("moyen_par_defaut", "Autre"),
                   sum(1 for o in toutes if o.moyen_paiement == cfg.data.get("moyen_par_defaut"))])
    _tableau(ws, ["Ordre", "Regex", "Moyen", "Nb opérations"], lignes, [8, 34, 20, 13],
             editables=("Ordre", "Regex", "Moyen"),
             note="La ligne « défaut » donne le moyen utilisé quand aucune regex ne correspond.")

    # Catégories banque -> type
    ws = wb.create_sheet("Catégories banque")
    lignes = [[k, v] for k, v in cfg.data.get("categories_banque", {}).items() if not k.startswith("_")]
    _tableau(ws, ["Catégorie banque", "Type"], lignes, [42, 18], editables=("Catégorie banque", "Type"),
             note="Utilisé quand aucune règle de libellé ne fixe le type.")
    _liste_types(ws, "B", len(lignes))

    # Comptes
    ws = wb.create_sheet("Comptes")
    lignes = [[num, nom] for num, nom in cfg.data.get("comptes", {}).items()]
    _tableau(ws, ["Numéro de compte", "Nom"], lignes, [20, 14], editables=("Numéro de compte", "Nom"),
             note="Numéro tel qu'il apparaît dans la colonne 'Numéro Compte' du CSV.")
    for r in range(2, len(lignes) + 2):
        ws.cell(r, 1).number_format = "@"


def importer_config_excel(chemin_xlsx: str | Path, cfg: Config, sauver: bool = True) -> Config:
    """
    Relit les onglets de configuration d'un classeur produit par ecrire_excel
    (Règles libellés, Contextes, Règles contextes, Types, Moyens paiement,
    Catégories banque, Comptes) et met à jour config_comptes.json.
    Une copie de l'ancien JSON est gardée en .bak.
    """
    from openpyxl import load_workbook
    wb = load_workbook(chemin_xlsx)

    def lignes(nom):
        if nom not in wb.sheetnames:
            return None
        ws = wb[nom]
        ent = [c.value for c in ws[1]]
        res = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            d = {e: v for e, v in zip(ent, row) if e}
            if not any(v not in (None, "") for v in d.values()):
                break                      # 1re ligne vide = fin du tableau (la note suit)
            res.append(d)
        return res

    txt = lambda v: "" if v is None else str(v).strip()
    brut = lambda v: "" if v is None else str(v)       # regex : on garde les espaces
    data = cfg.data

    if (l := lignes("Types")) is not None:
        data["types"] = [txt(x["Type"]) for x in l if txt(x.get("Type"))]
        data["types_hors_analyse"] = [txt(x["Type"]) for x in l
                                      if txt(x.get("Hors analyse")).lower() == "oui"]
    if (l := lignes("Règles libellés")) is not None:
        l = [x for x in l if txt(x.get("Regex"))]
        l.sort(key=lambda x: float(x["Ordre"]) if isinstance(x.get("Ordre"), (int, float)) else 1e9)
        for x in l:
            re.compile(txt(x["Regex"]))
        data["libelles"] = [{"regex": brut(x["Regex"]), "libelle": txt(x.get("Libellé")),
                             "type": txt(x.get("Type")) or None} for x in l]
    if (l := lignes("Moyens paiement")) is not None:
        regles = [x for x in l if txt(x.get("Ordre")) != "défaut" and txt(x.get("Regex"))]
        regles.sort(key=lambda x: float(x["Ordre"]) if isinstance(x.get("Ordre"), (int, float)) else 1e9)
        data["moyens_paiement"] = [{"regex": brut(x["Regex"]), "moyen": txt(x["Moyen"])} for x in regles]
        defaut = [x for x in l if txt(x.get("Ordre")) == "défaut"]
        if defaut:
            data["moyen_par_defaut"] = txt(defaut[0].get("Moyen")) or "Autre"
    if (l := lignes("Catégories banque")) is not None:
        aide = {k: v for k, v in data.get("categories_banque", {}).items() if k.startswith("_")}
        data["categories_banque"] = {**aide, **{txt(x["Catégorie banque"]): txt(x.get("Type"))
                                                for x in l if txt(x.get("Catégorie banque"))}}
    if (l := lignes("Comptes")) is not None:
        data["comptes"] = {txt(x["Numéro de compte"]): txt(x["Nom"]) for x in l
                           if txt(x.get("Numéro de compte"))}
        for nom in data["comptes"].values():
            if nom not in data["noms_comptes"]:
                data["noms_comptes"].append(nom)
    if (l := lignes("Contextes")) is not None:
        regles_ctx = lignes("Règles contextes") or []
        anciens = {c.get("nom"): c for c in data.get("contextes", [])}
        nouveaux = []
        for x in l:
            nom = txt(x.get("Nom"))
            if not nom:
                continue
            comptes = txt(x.get("Comptes"))
            remp = {}
            for morceau in txt(x.get("Types remplacés")).split(";"):
                if "→" in morceau or "->" in morceau:
                    a, b = re.split(r"→|->", morceau, maxsplit=1)
                    remp[a.strip()] = b.strip()
            nouveaux.append({**{k: v for k, v in anciens.get(nom, {}).items() if k.startswith("_")},
                             "nom": nom, "categorie": txt(x.get("Catégorie")),
                             "debut": _date(x["Début"]).isoformat(), "fin": _date(x["Fin"]).isoformat(),
                             "comptes": None if comptes.lower() in ("", "tous") else
                             [c.strip() for c in comptes.split(",") if c.strip()],
                             "regles": [{"regex": brut(r["Regex"]), "libelle": txt(r.get("Libellé")),
                                         "type": txt(r.get("Type")) or None}
                                        for r in regles_ctx if txt(r.get("Contexte")) == nom
                                        and txt(r.get("Regex"))],
                             "remplacer_types": remp,
                             "actif": txt(x.get("Actif")).lower() != "non"})
        data["contextes"] = nouveaux

    if sauver:
        bak = cfg.chemin.with_suffix(".json.bak")
        if cfg.chemin.exists():
            bak.write_text(cfg.chemin.read_text(encoding="utf-8"), encoding="utf-8")
        with open(cfg.chemin, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
    cfg._compiler()
    return cfg


# --------------------------------------------------------------------------- #
#  Ligne de commande
# --------------------------------------------------------------------------- #
def main() -> None:
    p = argparse.ArgumentParser(description="Relevé de comptes CSV -> Excel")
    p.add_argument("csv", nargs="*", help="fichier(s) CSV exportés de la banque")
    p.add_argument("-o", "--sortie", help="classeur Excel (défaut : Comptes.xlsx à côté de la config)")
    p.add_argument("-c", "--config", help="fichier de configuration (défaut : donnees\\config_comptes.json)")
    p.add_argument("--compte", help="force le nom du compte (Cyrille, Marion, Joint)")
    p.add_argument("--nouveau", action="store_true", help="écrase le classeur au lieu d'ajouter")
    p.add_argument("--afficher", action="store_true", help="affiche les opérations dans le terminal")
    p.add_argument("--analyse", metavar="TYPE", nargs="?", const="*",
                   help="affiche le coût par contexte pour un type (sans valeur : toutes dépenses)")
    p.add_argument("--importer-config", metavar="XLSX",
                   help="met à jour config_comptes.json depuis les onglets de configuration d'un classeur")
    a = p.parse_args()
    console_utf8()

    cfg = charger_config(a.config)
    a.sortie = a.sortie or str(cfg.chemin.parent / "Comptes.xlsx")
    if a.importer_config:
        importer_config_excel(a.importer_config, cfg)
        print(f"Configuration mise à jour depuis {a.importer_config} -> {cfg.chemin} "
              f"(ancienne version : {cfg.chemin.with_suffix('.json.bak').name})")
        if not a.csv:
            return
    if not a.csv:
        p.error("indiquez au moins un fichier CSV")
    a.csv = etendre_motifs(a.csv)
    ops: list[Operation] = []
    for f in a.csv:
        ops += lire_csv(f, cfg, compte=a.compte)
    calculer_totaux(ops, cfg)
    if a.afficher:
        for o in ops:
            print(o)
    a_classer = [o for o in ops if o.type == cfg.data.get("type_par_defaut")]
    if a_classer:
        print(f"{len(a_classer)} opération(s) 'À classer' : ajoutez une règle dans 'libelles' de la config.")
    if a.analyse:
        t = None if a.analyse == "*" else a.analyse
        afficher_analyse(analyser_contextes(ops, cfg, t), f"Coût par contexte – {t or 'toutes dépenses'}")
    ecrire_excel(ops, a.sortie, cfg, ajouter=not a.nouveau)


if __name__ == "__main__":
    main()
