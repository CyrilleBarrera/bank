"""Tests sur des données FICTIVES (tests/donnees/export_fictif.csv). Lancer : pytest"""
import json
from datetime import date
from pathlib import Path

import pytest

from comptes_bancaires import (QUOTIDIEN, Config, analyser_contextes, calculer_totaux, ecrire_excel,
                               importer_config_excel, lire_csv)
from comptes_bancaires.releve import _lire_existant, _nombre, date_achat, etendre_motifs

ICI = Path(__file__).parent
CSV = ICI / "donnees" / "export_fictif.csv"
MODELE = ICI.parent / "src" / "comptes_bancaires" / "config_defaut.json"


@pytest.fixture
def cfg(tmp_path) -> Config:
    """Config modèle + compte fictif + un contexte « Vacances test » du 15 au 19 mars."""
    data = json.loads(MODELE.read_text(encoding="utf-8"))
    data["comptes"] = {"00011112222": "Cyrille"}
    data["contextes"] = [{"nom": "Vacances test", "categorie": "Vacances", "debut": "2026-03-15",
                          "fin": "2026-03-19", "comptes": None, "regles": [],
                          "remplacer_types": {"Restaurant": "Voyage"}, "actif": True}]
    chemin = tmp_path / "config_comptes.json"
    chemin.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    return Config.charger(chemin)


@pytest.fixture
def ops(cfg):
    o = lire_csv(CSV, cfg)
    calculer_totaux(o, cfg)
    return o


def test_nombres_francais():
    assert _nombre("-1 800,00") == -1800.0
    assert _nombre("2 736,36") == 2736.36
    assert _nombre("413.98") == 413.98
    assert _nombre("") == 0.0


def test_date_achat_carte():
    assert date_achat("CARTE 18/03/26 BRASSERIE CB*1234", date(2026, 3, 20)) == date(2026, 3, 18)
    assert date_achat("PRLV SEPA SFR", date(2026, 3, 12)) == date(2026, 3, 12)


def test_lecture_csv(ops):
    assert len(ops) == 9
    assert ops[0].date == date(2026, 2, 28)                    # trié du plus ancien au plus récent
    assert ops[0].date_texte == "28 Février 2026"
    inter = [o for o in ops if o.libelle == "Intermarché"]
    assert len(inter) == 2 and all(o.type == "Nourriture" and o.debit > 0 and o.credit == 0 for o in inter)
    assert {o.moyen_paiement for o in ops} >= {"Carte Bleue", "Virement", "Prélèvement"}
    sfr = next(o for o in ops if o.libelle == "Abonnement SFR")
    assert sfr.type == "Charges Fixes" and sfr.commentaire == "test commentaire"
    salaire = next(o for o in ops if "STEERLIGHT" in o.libelle_banque)
    assert (salaire.libelle, salaire.type) == ("Salaire Cyrille", "Revenus")
    marion = next(o for o in ops if "MAINSANT" in o.libelle_banque)
    assert (marion.libelle, marion.type) == ("Equilibrage comptes", "Marion")
    assert next(o for o in ops if "CPAM" in o.libelle_banque.replace(".", "")).libelle == "Remboursement santé"


def test_totaux_egal_solde_banque(ops):
    # le dernier total recalculé doit retomber sur le solde affiché par la banque
    assert ops[-1].total_cyrille == 1830.0
    assert ops[-1].total_general == 1830.0
    assert ops[-1].total_marion == 0.0


def test_contexte_date_achat_et_remplacement(ops):
    brasserie = next(o for o in ops if "BRASSERIE" in o.libelle_banque)
    assert brasserie.contexte == "Vacances test"               # achat le 18, débit le 20
    assert brasserie.type == "Voyage"                           # Restaurant -> Voyage pendant le contexte
    inter_31 = next(o for o in ops if o.date == date(2026, 3, 31))
    assert inter_31.contexte == QUOTIDIEN


def test_analyse_contextes(ops, cfg):
    lignes = {l["contexte"]: l for l in analyser_contextes(ops, cfg, "Nourriture")}
    assert lignes["Vacances test"]["jours"] == 5
    assert lignes["Vacances test"]["depense"] == 30.0          # Intermarché du 17/03
    assert lignes[QUOTIDIEN]["depense"] == 20.0


def test_excel_ajout_sans_doublon_et_corrections_gardees(ops, cfg, tmp_path):
    xlsx = tmp_path / "Comptes.xlsx"
    ecrire_excel(ops, xlsx, cfg, ajouter=False)
    # correction manuelle simulée dans Excel
    from openpyxl import load_workbook
    wb = load_workbook(xlsx)
    ws = wb["Opérations"]
    entete = [c.value for c in ws[1]]
    ws.cell(2, entete.index("Commentaire") + 1, "corrigé à la main")
    wb.save(xlsx)
    ecrire_excel(ops, xlsx, cfg)                                # même CSV une 2e fois
    relues, soldes = _lire_existant(xlsx, cfg)
    assert len(relues) == 9
    assert relues[0].commentaire == "corrigé à la main"
    assert soldes["Cyrille"] == 500.0
    attendus = {"Opérations", "Synthèse", "Évolution", "Règles libellés", "Contextes", "Types"}
    assert attendus <= set(load_workbook(xlsx).sheetnames)


def test_import_config_depuis_excel(ops, cfg, tmp_path):
    xlsx = tmp_path / "Comptes.xlsx"
    ecrire_excel(ops, xlsx, cfg, ajouter=False)
    avant = json.loads(cfg.chemin.read_text(encoding="utf-8"))
    importer_config_excel(xlsx, cfg)
    apres = json.loads(cfg.chemin.read_text(encoding="utf-8"))
    for cle in ("libelles", "moyens_paiement", "comptes", "types"):
        assert avant[cle] == apres[cle], cle
    assert set(avant["types_hors_analyse"]) == set(apres["types_hors_analyse"])
    assert cfg.chemin.with_suffix(".json.bak").exists()


def test_graphiques_html(ops, cfg, tmp_path):
    from comptes_bancaires.graphiques import generer_html
    sortie = generer_html(ops, cfg, tmp_path / "g.html", periode="semaine")
    texte = sortie.read_text(encoding="utf-8")
    assert "Évolution d'une catégorie" in texte and "plotly" in texte.lower()


def test_motifs_windows(tmp_path):
    for n in ("a.csv", "b.csv", "c.txt"):
        (tmp_path / n).write_text("x")
    assert len(etendre_motifs([str(tmp_path / "*.csv")])) == 2
    with pytest.raises(SystemExit):
        etendre_motifs([str(tmp_path / "*.xls")])


# --------------------------------------------------------------------------- #
#  Règles v2 (rapprochement avec le classeur tenu à la main)
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize("brut, libelle, type_", [
    ("CARTE 28/09/26 CITY CB*7901", "Carrefour City", "Nourriture"),
    ("CARTE 22/09/26 SAS ENZO CB*7901", "Estacad : Légumes", "Nourriture"),
    ("CARTE 03/10/26 COSTA CB*7901", "Estacad : Légumes", "Nourriture"),
    ("CARTE 03/10/26 BIOTYFULL FRUIT CB*7901", "Estacad : Fruits", "Nourriture"),
    ("CARTE 16/09/26 COMPASS-6505 2 CB*7901", "Cantine CEA", "Nourriture"),
    ("RETRAIT DAB 25/09/26 ST MARTIN D CB*7901", "Retrait liquide", "Vie courante"),
    ("CARTE 26/08/26 SUMUP *PIZZA M7R CB*7901", "Fast Pizza", "Restaurant"),
    ("VIR LDDS", "Remplissage LDDS", "LivretA"),
    ("VIR INST PATRICK PHILIP", "Covoit CAFGO", "Loisirs"),
    ("CARTE 25/09/26 INTERM CALAO 76 9 CB*7901", "Intermarché", "Nourriture"),
    ("CARTE 19/08/26 SNCF-VOYAGEURS CB*7901", "SNCF", "Transport"),
])
def test_regles_v2(cfg, brut, libelle, type_):
    assert cfg.libelle_et_type(brut, "", "", "") == (libelle, type_)


def test_virements_variables_a_classer(cfg):
    # les virements dont l'objet change à chaque fois ne reçoivent pas de type figé
    assert cfg.libelle_et_type("VIR Virement de LUIS BARRERA", "", "Virements reçus", "") == \
        ("Virement Luis Barrera", "À classer")


def test_restaurant_par_categorie(cfg):
    assert cfg.libelle_et_type("CARTE 01/01/26 NOUVEAU RESTO CB*1", "Nouveau Resto",
                               "Restaurants, bars, discothèques…", "Loisirs et sorties")[1] == "Restaurant"


def test_avoir_carte(cfg):
    assert cfg.moyen_paiement("AVOIR 18/01/26 WWW.SKITRUCK.FR CB*7901") == "Carte Bleue"


# --------------------------------------------------------------------------- #
#  Migration d'une configuration v1 (installée par la version 0.4)
# --------------------------------------------------------------------------- #
def test_migration_garde_les_reglages_perso(tmp_path):
    from comptes_bancaires.migration import migrer_config
    ref = json.loads((MODELE.parent / "reference_config_v1.json").read_text(encoding="utf-8"))
    v1 = json.loads(MODELE.read_text(encoding="utf-8"))
    v1.pop("version_config")
    v1["libelles"] = [{"regex": "MON MAGASIN", "libelle": "Mon magasin", "type": "Choeur"}] + ref["libelles"]
    v1["types"] = ref["types"] + ["Bricolage"]
    v1["comptes"] = {"123": "Marion"}
    v1["categories_banque"] = dict(ref["categories_banque"], **{"Alimentation": "Courses"})
    chemin = tmp_path / "config_comptes.json"
    chemin.write_text(json.dumps(v1, ensure_ascii=False), encoding="utf-8")

    journal = migrer_config(chemin)
    data = json.loads(chemin.read_text(encoding="utf-8"))
    assert journal and data["version_config"] == 2
    assert data["libelles"][0] == {"regex": "MON MAGASIN", "libelle": "Mon magasin", "type": "Chœurs"}
    assert any(r["libelle"] == "Carrefour City" for r in data["libelles"])
    assert not any(r["regex"] == "VIR SEPA STEERLIGHT" for r in data["libelles"])   # ancienne règle remplacée
    assert {"Restaurant", "Revenus", "Bricolage", "Chœurs"} <= set(data["types"]) and "Choeur" not in data["types"]
    assert data["types"][-1] == "À classer"
    assert data["comptes"] == {"123": "Marion"}
    assert data["categories_banque"]["Alimentation"] == "Courses"                       # choix perso gardé
    assert data["categories_banque"]["Restaurants, bars, discothèques…"] == "Restaurant"
    assert (tmp_path / "config_comptes.avant-v2.json").exists()
    assert migrer_config(chemin) == []                                                   # idempotent
    Config.charger(chemin)                                                               # reste valide


def test_lanceur_relance_dans_le_venv():
    import _lanceur  # noqa: F401  (le module se charge sans effet de bord)
    assert _lanceur.RACINE.joinpath("comptes.py").exists()
